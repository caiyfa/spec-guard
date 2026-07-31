"""file_preview 办公文档转换器测试（K-030）。

验证：
1. xlsx → Markdown 表格（表头/数据行/截断标记/空行跳过）
2. docx → 纯文本（段落 + 表格提取）
3. pdf → 错误路径（损坏文件不抛异常，返回明确文本）
4. read_spec 对 xlsx 返回转换文本而非拒读
"""

import sys
from pathlib import Path

import pytest

SG_ROOT = Path(__file__).resolve().parent.parent
if str(SG_ROOT) not in sys.path:
    sys.path.insert(0, str(SG_ROOT))
BACKEND_DIR = SG_ROOT / "workbench" / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from services import file_preview


class TestXlsxPreview:
    def test_xlsx_to_markdown_table(self, tmp_path):
        """xlsx 转换为 Markdown 表格（含 sheet 标题与表头）。"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "用例"
        ws.append(["ID", "名称", "优先级"])
        ws.append(["TC-001", "创建标签", "P1"])
        ws.append([None, None, None])  # 全空行应跳过
        ws.append(["TC-002", "删除|标签", "P2"])  # 竖线应转义
        f = tmp_path / "cases.xlsx"
        wb.save(f)

        fmt, text = file_preview.preview_text(f)
        assert fmt == "markdown"
        assert "## Sheet: 用例" in text
        assert "| ID | 名称 | 优先级 |" in text
        assert "| TC-001 | 创建标签 | P1 |" in text
        assert "删除\\|标签" in text  # 竖线转义
        assert text.count("TC-00") == 2  # 空行未计入

    def test_xlsx_truncation_marker(self, tmp_path):
        """超 200 行的 sheet 附截断标记。"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["ID"])
        for i in range(300):
            ws.append([f"TC-{i:04d}"])
        f = tmp_path / "big.xlsx"
        wb.save(f)
        _fmt, text = file_preview.preview_text(f)
        assert "已截断至 200 行" in text

    def test_xlsx_corrupt_no_crash(self, tmp_path):
        """损坏文件返回明确错误文本（不抛异常）。"""
        f = tmp_path / "bad.xlsx"
        f.write_bytes(b"not a real xlsx")
        _fmt, text = file_preview.preview_text(f)
        assert "解析失败" in text

    def test_xlsx_merged_cells_filled(self, tmp_path):
        """K-035：合并单元格（多级表头）回填左上角值，行列对齐不错位。"""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["业务单元", "姓名", "打标A", None, "打标B", None])
        ws.append([None, None, "总分", "排名", "总分", "排名"])
        ws.append(["远东", "程洋", "98", "1/5", "97", "2/5"])
        ws.merge_cells("A1:A2")
        ws.merge_cells("B1:B2")
        ws.merge_cells("C1:D1")
        ws.merge_cells("E1:F1")
        f = tmp_path / "merged.xlsx"
        wb.save(f)

        _fmt, text = file_preview.preview_text(f)
        # 合并表头回填：打标A 占两列
        assert "| 业务单元 | 姓名 | 打标A | 打标A | 打标B | 打标B |" in text
        # 纵向合并回填：第二行首列不再是空白
        assert "| 业务单元 | 姓名 | 总分 | 排名 | 总分 | 排名 |" in text
        assert "| 远东 | 程洋 | 98 | 1/5 | 97 | 2/5 |" in text

    def test_xlsx_multi_sheet_all_rendered(self, tmp_path):
        """K-035：多 sheet 全部渲染（各自的 sheet 标题 + 数据）。"""
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "总分"
        wb.active.append(["列A"])
        wb.active.append(["数据A"])
        ws2 = wb.create_sheet("能力")
        ws2.append(["列B"])
        ws2.append(["数据B"])
        f = tmp_path / "multi.xlsx"
        wb.save(f)

        _fmt, text = file_preview.preview_text(f)
        assert "## Sheet: 总分" in text and "数据A" in text
        assert "## Sheet: 能力" in text and "数据B" in text

    def test_xlsx_bad_dimension_still_reads_all(self, tmp_path):
        """K-035：sheet XML dimension 声明错误（非 Excel 导出文件常见）时仍读全表。

        read_only 模式信任 dimension（坏声明 → 只读 1 行 1 列，线上事故根因），
        普通模式按真实单元格解析，免疫坏 dimension。
        """
        import zipfile, shutil
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["H1", "H2", "H3"])
        ws.append(["a", "b", "c"])
        f = tmp_path / "bad_dim.xlsx"
        wb.save(f)
        # 篡改 dimension 声明为 A1:A1（模拟问题文件）
        fixed = tmp_path / "fixed.xlsx"
        with zipfile.ZipFile(f) as zin, zipfile.ZipFile(fixed, "w") as zout:
            for item in zin.namelist():
                data = zin.read(item)
                if item == "xl/worksheets/sheet1.xml":
                    import re
                    data = re.sub(rb'<dimension ref="[^"]*"/>', rb'<dimension ref="A1:A1"/>', data)
                zout.writestr(item, data)

        _fmt, text = file_preview.preview_text(fixed)
        assert "| H1 | H2 | H3 |" in text
        assert "| a | b | c |" in text


class TestDocxPreview:
    def test_docx_to_text(self, tmp_path):
        """docx 提取段落与表格文本。"""
        import docx
        doc = docx.Document()
        doc.add_paragraph("第一段：标签库维护需求")
        doc.add_paragraph("")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "类型"
        table.cell(1, 0).text = "tagName"
        table.cell(1, 1).text = "string"
        f = tmp_path / "spec.docx"
        doc.save(f)

        fmt, text = file_preview.preview_text(f)
        assert fmt == "text"
        assert "第一段：标签库维护需求" in text
        assert "tagName | string" in text


class TestPdfPreview:
    def test_pdf_corrupt_no_crash(self, tmp_path):
        """损坏文件返回明确错误文本。"""
        f = tmp_path / "bad.pdf"
        f.write_bytes(b"%PDF-1.4 broken")
        _fmt, text = file_preview.preview_text(f)
        assert "解析失败" in text or "无可提取文本" in text


class TestReadSpecOfficeConversion:
    """read_spec 对办公文档返回转换文本（K-030 前是拒读）。"""

    @pytest.fixture()
    def office_ws(self, tmp_path):
        from specstd.workspace import register_local, unregister
        from test_specstd.conftest import create_valid_workspace
        ws_path = create_valid_workspace(tmp_path / "ws")
        register_local("test-preview", str(ws_path), "tdm", activate=True)
        yield ws_path
        unregister("test-preview")

    def test_read_spec_xlsx_converted(self, office_ws):
        import asyncio
        from openpyxl import Workbook
        from mcp_server.tools import spec_reader

        xlsx_path = office_ws / "specs" / "module1" / "feature1" / "cases.xlsx"
        wb = Workbook()
        wb.active.append(["ID", "名称"])
        wb.active.append(["TC-100", "转换验证"])
        wb.save(xlsx_path)

        result = asyncio.run(spec_reader.handle(
            "read_spec", {"path": "specs/module1/feature1/cases.xlsx"}))
        text = result[0].text
        assert "已转换为文本预览" in text
        assert "TC-100" in text and "转换验证" in text
        assert "二进制文件不可文本读取" not in text

    def test_read_file_raw_roundtrip(self, office_ws):
        """K-036：read-file?raw=true 返回原始字节（base64）——保真预览的数据通道。

        raw=true → 字节级 roundtrip（SheetJS 前端解析需要原文件）；
        raw=false → 仍返回转换文本（Agent 读取链路不受影响）。
        """
        import base64
        from openpyxl import Workbook
        from routers.workspace import api_read_file

        xlsx_path = office_ws / "specs" / "module1" / "feature1" / "raw.xlsx"
        wb = Workbook()
        wb.active.append(["ID"])
        wb.save(xlsx_path)
        original = xlsx_path.read_bytes()

        r = api_read_file("specs/module1/feature1/raw.xlsx", raw=True, ws_name="test-preview")
        assert r["format"] == "xlsx"
        assert base64.b64decode(r["raw_base64"]) == original  # 字节级一致

        r2 = api_read_file("specs/module1/feature1/raw.xlsx", raw=False, ws_name="test-preview")
        assert r2["format"] == "markdown" and "raw_base64" not in r2
