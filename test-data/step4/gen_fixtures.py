# -*- coding: utf-8 -*-
"""生成 SpecGuard 导入场景测试 fixture（需求驱动，不读实现改预期）。

fixture 与场景的对应关系：
- eval360-empty-scores.xlsx     → EV-IMP-001（360分数 "-"/空值 可导入且不丢其余数据）
- eval360-required-missing.xlsx → EV-IMP-002（必填校验保留 + 行级隔离）
- grid-position-valid.xlsx      → GP-IMP-001（落位 "0" 与 "-" 视为没有分数）
- grid-position-out-of-range.xlsx → GP-IMP-002（落位越界报错 + 行级隔离）

列序严格对齐 DTO @ExcelProperty index（EasyExcel 按 index 读取，表头仅装饰）：
- 公共列 0-5：工号/姓名/岗位/部门全路径名称/部门全路径编码/高潜标记
"""
import os
from openpyxl import Workbook

BASE = ["工号", "姓名", "岗位", "部门全路径名称", "部门全路径编码", "高潜标记"]

HEAD_360_TOTAL = BASE + ["套卷名称", "类型", "总分", "总人数", "总排名", "套卷内总人数",
                         "套卷内排名", "管理层排名", "关键纵向排名",
                         "自评分", "上级评分", "下级评分", "合作伙伴评分"]
HEAD_360_DIM = BASE + ["维度类别", "套卷名称", "维度得分", "总人数", "总排名", "套卷内总人数",
                       "套卷排名", "管理层排名", "关键纵向排名",
                       "自评分", "上级评分", "下级评分", "合作伙伴评分"]
HEAD_360_VALUE = HEAD_360_DIM[:6] + ["套卷名称", "行为文案"] + HEAD_360_DIM[8:]
HEAD_EVAL = BASE + ["分管领导校准后落位", "分管领导评价等级", "人力独立评价等级",
                    "履职状态等级", "优势", "待提升", "任用建议", "风险事项"]

# 三个测试工号（需在 DHR 盘点日期 2026-07-22 存在；PREP 场景会加入会话人员清单）
EMP1 = ["90024713", "张测试一", "总经理", "远东国际/总经办", "50000001", "否"]
EMP2 = ["90027339", "张测试二", "总监", "远东国际/业务部", "50000002", "否"]
EMP3 = ["90027338", "张测试三", "经理", "远东国际/职能部", "50000003", "否"]

OUT_DIR = os.path.dirname(os.path.abspath(__file__))


def pad(row, width):
    return row + [None] * (width - len(row))


def gen_eval360_empty_scores():
    wb = Workbook()

    ws = wb.active
    ws.title = "360评价总分"
    ws.append(HEAD_360_TOTAL)
    # r1: 全部分数列填 "-" → 分数入库应为 null，但总人数/总排名等非分数字段同样为 "-" → null
    ws.append(pad(EMP1 + ["SG-套卷A", "总分", "-", "-", "-", "-", "-", "-", "-", "-", "-", "-", "-"], len(HEAD_360_TOTAL)))
    # r2: 分数列全部留空 → 入库为 null（空值可导入）
    ws.append(pad(EMP2 + ["SG-套卷B", "总分"], len(HEAD_360_TOTAL)))
    # r3: 正常分数行（对照组），验证合法值不受影响
    ws.append(pad(EMP3 + ["SG-套卷C", "总分", "88.5", "20", "3", "10", "2", "5", "4", "90", "88", "87", "89"], len(HEAD_360_TOTAL)))

    ws = wb.create_sheet("360维度得分")
    ws.append(HEAD_360_DIM)
    # r1: 维度得分 "-" → null，总排名=5 保留（不要删掉其余数据）
    ws.append(pad(EMP1 + ["决策能力", "SG-套卷A", "-", "10", "5", None, None, None, None, None, None, None, None], len(HEAD_360_DIM)))
    # r2/r3: 正常行（回归："-" 不得导致整个 Sheet 解析失败丢数据）
    ws.append(pad(EMP1 + ["沟通能力", "SG-套卷A", "92.5"], len(HEAD_360_DIM)))
    ws.append(pad(EMP2 + ["决策能力", "SG-套卷B", "85"], len(HEAD_360_DIM)))

    ws = wb.create_sheet("360价值观")
    ws.append(HEAD_360_VALUE)
    ws.append(pad(EMP1 + ["SG-套卷A", "诚信正直", "-"], len(HEAD_360_VALUE)))
    ws.append(pad(EMP2 + ["SG-套卷B", "客户至上", "95"], len(HEAD_360_VALUE)))

    path = os.path.join(OUT_DIR, "eval360-empty-scores.xlsx")
    wb.save(path)
    return path


def gen_eval360_required_missing():
    wb = Workbook()
    ws = wb.active
    ws.title = "360评价总分"
    ws.append(HEAD_360_TOTAL)
    # r1: 类型留空 → 必填报错（[360评价总分] 以下字段必填：类型）
    ws.append(pad(EMP1 + ["SG-套卷D", None, "80"], len(HEAD_360_TOTAL)))
    # r2: 正常行 → 成功（行级隔离：r1 报错不影响 r2 入库）
    ws.append(pad(EMP2 + ["SG-套卷E", "总分", "80"], len(HEAD_360_TOTAL)))
    path = os.path.join(OUT_DIR, "eval360-required-missing.xlsx")
    wb.save(path)
    return path


def _eval_row(emp, placement):
    return pad(emp + [placement, "优秀", "良好", "正常", "逻辑思维强", "沟通能力待提升", "可任用", "无"], len(HEAD_EVAL))


def gen_grid_position_valid():
    wb = Workbook()
    ws = wb.active
    ws.title = "综合评价"
    ws.append(HEAD_EVAL)
    ws.append(_eval_row(EMP1, "0"))    # 落位 0（没有分数）→ 合法入库
    ws.append(_eval_row(EMP2, "-"))    # 落位 "-"（没有分数）→ 合法入库，DB 存 null
    ws.append(_eval_row(EMP3, "5"))    # 对照组：正常落位
    path = os.path.join(OUT_DIR, "grid-position-valid.xlsx")
    wb.save(path)
    return path


def gen_grid_position_out_of_range():
    wb = Workbook()
    ws = wb.active
    ws.title = "综合评价"
    ws.append(HEAD_EVAL)
    ws.append(_eval_row(EMP1, "10"))   # 越界 → 行报错「必须为0-9的数字」
    ws.append(_eval_row(EMP2, "7"))    # 正常行 → 成功（行级隔离）
    path = os.path.join(OUT_DIR, "grid-position-out-of-range.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    for p in (gen_eval360_empty_scores(), gen_eval360_required_missing(),
              gen_grid_position_valid(), gen_grid_position_out_of_range()):
        print("generated:", p)
